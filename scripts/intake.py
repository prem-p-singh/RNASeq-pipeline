"""Read the existing questionnaire workbook into setup inputs (M1 first slice).

No Excel formulas or macros are executed. The legacy template is matched by
question label, never row number; future templates may carry field IDs in E.
"""
from pathlib import Path
import hashlib
import math
import re

import yaml
import metadata

ROOT = Path(__file__).resolve().parent.parent


def read_intake(path, sheet_name=None):
    import openpyxl

    path = Path(path).expanduser().resolve()
    definition = yaml.safe_load((ROOT / "config/intake_questions.yaml").read_text())
    tabs = {tab["name"]: tab for tab in definition["tabs"]}
    workbook = openpyxl.load_workbook(path, data_only=False, keep_links=False)
    try:
        version = workbook["README"]["B2"].value if "README" in workbook.sheetnames else None
        version = 1 if version in (None, "") else version
        if type(version) is not int or version not in (1, 2, 3):
            raise ValueError("README!B2: unsupported intake schema version; expected 1, 2 or 3")
        candidates = []
        for name in tabs:
            if name in workbook.sheetnames:
                # Defaults exist on every tab. A project name selects a tab.
                if any((row[4].value == "project_name" or row[1].value == "Project name")
                       and row[2].value not in (None, "")
                       for row in workbook[name].iter_rows(min_col=1, max_col=5)):
                    candidates.append(name)
        if sheet_name is None:
            if len(candidates) != 1:
                raise ValueError("Fill Project name on exactly one assay tab, or select --intake-sheet explicitly.")
            sheet_name = candidates[0]
        if sheet_name not in tabs or sheet_name not in workbook.sheetnames:
            raise ValueError(f"Unknown intake tab: {sheet_name!r}")
        questions = definition["common_questions"] + tabs[sheet_name]["extra_questions"]
        by_label = {q["text"]: q for q in questions}
        by_id = {q["id"]: q for q in questions}
        answers, sources, errors = {}, {}, []
        for row in workbook[sheet_name].iter_rows(min_col=1, max_col=5):
            label, cell, field = row[1].value, row[2], row[4].value
            # Accept prior site-specific labels without binding to a host name.
            if isinstance(label, str):
                label = re.sub(r"^FASTQ folder path \(on [^)]+\)$", "FASTQ folder path (on analysis host)", label)
                label = re.sub(r"^(Metadata file \(CSV or Excel\) — full path on ).+$", r"\1analysis host", label)
            question = by_id.get(field) if field else by_label.get(label)
            location = f"{sheet_name}!{cell.coordinate}"
            if question is None:
                if cell.value not in (None, "", "Your answer"):
                    errors.append(f"{location}: unrecognized question/field {field or label!r}")
                continue
            key = question["id"]
            if key in answers:
                errors.append(f"{location}: duplicate field {key}")
                continue
            value = cell.value
            if cell.data_type == "f":
                errors.append(f"{location}: enter a literal answer, not an Excel formula")
                continue
            if isinstance(value, str):
                value = value.strip()
            if value in (None, ""):
                if question.get("required") and key not in ("fastq_path", "metadata_file", "sample_id_column"):
                    errors.append(f"{location}: {question['text']} is required")
                value = None
            elif question["type"] == "dropdown" and value not in question["options"]:
                errors.append(f"{location}: choose one of {question['options']}")
            elif question["type"] == "number":
                if isinstance(value, bool) or not isinstance(value, (int, float)) or not math.isfinite(value) or value < 1 or int(value) != value:
                    errors.append(f"{location}: expected a positive whole number")
            elif not isinstance(value, (str, int, float)) or isinstance(value, bool):
                errors.append(f"{location}: expected text or a number; dates are not identifiers")
            answers[key] = value
            sources[key] = location
        for q in questions:
            if q.get("required") and q.get("introduced_in", 1) <= version and q["id"] not in answers:
                errors.append(f"{sheet_name}: missing required question {q['id']}")
        inline = answers.get("metadata_source") == "workbook tables"
        table_records = {}
        for name in metadata.SCHEMAS:
            title = name.title()
            if title not in workbook.sheetnames:
                if inline:
                    errors.append(f"Missing worksheet {title}")
                continue
            ws = workbook[title]
            headers = [c.value for c in ws[4]]
            while headers and headers[-1] is None:
                headers.pop()
            if not headers or any(not isinstance(c, str) or not c.strip() for c in headers) or len(set(headers)) != len(headers):
                errors.append(f"{title}!4: expected unique nonblank column names")
                continue
            rows = []
            for row in ws.iter_rows(min_row=5):
                if all(c.value in (None, "") for c in row):
                    continue
                values = {}
                for i, cell in enumerate(row):
                    if i >= len(headers):
                        if cell.value not in (None, ""):
                            errors.append(f"{title}!{cell.coordinate}: value has no column heading")
                        continue
                    key, value = headers[i], cell.value
                    location = f"{title}!{cell.coordinate}"
                    if cell.data_type == "f" or isinstance(value, bool) or (value is not None and not isinstance(value, (str, int, float))):
                        errors.append(f"{location}: enter literal text or a number, not formulas or dates")
                    if key.endswith("_id") and value is not None and not isinstance(value, str):
                        errors.append(f"{location}: identifiers must be stored as Text before entry")
                    values[key] = str(value).strip() if value is not None else ""
                    sources[f"{name}.{cell.row}.{key}"] = location
                rows.append(values)
            table_records[name] = rows
        if inline:
            if answers.get("fastq_path") or answers.get("metadata_file"):
                errors.append("Workbook tables selected: clear FASTQ folder and external metadata file to avoid conflicting sources")
            errors.extend(f"{code}: {detail}" for code, detail in metadata.validate(table_records))
        else:
            if any(table_records.values()):
                errors.append("Workbook tables contain records: select workbook tables as metadata source or clear those records")
            for key in ("fastq_path", "metadata_file", "sample_id_column"):
                if not answers.get(key):
                    errors.append(f"{sources.get(key, sheet_name)}: {key} is required with external files")
        if errors:
            raise ValueError("Intake needs correction:\n  " + "\n  ".join(errors))
        return {
            "schema_version": version, "workbook_sha256": hashlib.sha256(path.read_bytes()).hexdigest(),
            "source": str(path), "sheet": sheet_name, "answers": answers, "sources": sources,
            "metadata_tables": table_records if inline else None,
        }
    finally:
        workbook.close()


def setup_inputs(record):
    """Convert only executable questionnaire choices; never invent a model."""
    a, sources = record["answers"], record["sources"]

    def fail(key, detail):
        raise ValueError(f"{sources.get(key, record['sheet'])}: {detail}")

    if record["sheet"] != "Bulk RNA-seq":
        fail("project_name", "Direct intake execution currently supports Bulk RNA-seq only. "
             "TAG-seq needs kit qualification; small RNA needs its own processing route. "
             "No bulk fallback was selected.")
    if a.get("input_stage", "raw FASTQ") != "raw FASTQ":
        fail("input_stage", "Only raw FASTQ input is implemented; no counts/object or instrument-data import is available.")
    if a.get("umi", "no") != "no":
        fail("umi", "Confirm a non-UMI library. UMI processing is not implemented; unknown status must be resolved.")
    if not re.fullmatch(r"[A-Za-z0-9_]+", str(a["project_name"])):
        fail("project_name", "use letters, digits and underscores")
    layout = {"paired-end": "rnaseq_paired", "single-end": "rnaseq_single"}.get(a["read_config"])
    if layout is None:
        fail("read_config", "read configuration must be resolved before setup")
    if a["repeated_measures"] == "yes":
        fail("repeated_measures", "This workbook has no subject/model fields yet. "
             "Use the YAML setup input with explicit fixed/random effects for repeated measures.")
    factor = str(a["primary_factor"])
    batch = a["batch_effect_column"]
    terms = [str(batch), factor] if batch and batch != "none" else [factor]
    if any(not re.fullmatch(r"[A-Za-z][A-Za-z0-9_]*", term) for term in terms):
        fail("primary_factor", "factor/batch column names must be simple identifiers")
    organism = str(a["organism"])
    if organism.isdigit():
        tax_id = int(organism)
    elif isinstance(a["organism"], (int, float)) and math.isfinite(a["organism"]) and int(a["organism"]) == a["organism"]:
        tax_id = int(a["organism"])
    else:
        matches = []
        for p in (ROOT / "scripts/presets").glob("*.yaml"):
            preset = yaml.safe_load(p.read_text())
            if organism.casefold() in {str(preset.get(k, "")).casefold() for k in ("scientific_name", "common_name")}:
                matches.append(preset["tax_id"])
        if len(set(matches)) != 1:
            fail("organism", "supply an NCBI taxID; this name has no unambiguous local preset")
        tax_id = matches[0]
    if tax_id <= 0:
        fail("organism", "NCBI taxID must be positive")

    def absolute(value):
        if "://" in value:
            return value
        p = Path(value).expanduser()
        return str(p if p.is_absolute() else Path(record["source"]).parent / p)

    strand = a["strandedness"]
    expected = {"reverse": "ISR" if layout == "rnaseq_paired" else "SR",
                "forward": "ISF" if layout == "rnaseq_paired" else "SF",
                "unstranded": "IU" if layout == "rnaseq_paired" else "U"}.get(strand)
    supplied = a.get("expected_libtype")
    if supplied and expected and supplied.upper() != expected:
        fail("expected_libtype", f"conflicts with declared strand/layout ({expected})")
    if supplied:
        supplied = supplied.upper()
        allowed = {"IU", "ISF", "ISR", "OU", "OSF", "OSR", "MU", "MSF", "MSR"} if layout == "rnaseq_paired" else {"U", "SF", "SR"}
        if supplied not in allowed:
            fail("expected_libtype", f"invalid library type for {a['read_config']}")
    if a["orgdb_strategy"] == "use_existing":
        fail("orgdb_strategy", "This workbook has no OrgDb package field. Use YAML with an explicit package.")
    enrichment = a["run_enrichment"] == "yes"
    if enrichment and a["orgdb_strategy"] == "skip":
        fail("orgdb_strategy", "skip conflicts with requested GO + KEGG enrichment; select no enrichment or a strategy")
    record["context_only"] = ["contact_email", "tissue_type", "platform", "library_type", "library_kit"]
    record["limitations"] = [
        "Platform and library preparation are recorded context; the existing bulk processor still uses its release preprocessing defaults.",
        "Contact email is recorded but does not configure scheduler notifications.",
        "Multiple libraries per sample, repeated-measures and non-bulk workbook execution remain pending assay qualification.",
    ]
    if record["schema_version"] == 1:
        record["limitations"].append("Legacy workbook: assumes raw non-UMI bulk input. Confirm protocol or use the current template with explicit eligibility fields.")
    tables = record.get("metadata_tables")
    if tables:
        units = [row.get("biological_unit_id") or row["sample_id"] for row in tables["samples"]]
        if len(set(units)) != len(units):
            fail("repeated_measures", "Repeated biological units require an explicit subject model; independent bulk intake cannot treat them as replicates")
        for read in tables["reads"]:
            read["uri"] = absolute(read["uri"])
    return dict(project_name=a["project_name"], tax_id=tax_id, seq_type=layout,
                metadata_tables=tables,
                fastq_source=absolute(str(a["fastq_path"])) if not tables else None,
                metadata_file=absolute(str(a["metadata_file"])) if not tables else None,
                sample_id_column=str(a["sample_id_column"]), primary_factor=factor,
                model_fixed_effects="~ " + " + ".join(dict.fromkeys(terms)), model_random_effects=None,
                expected_libtype=supplied or expected, description=a["biological_question"],
                run_enrichment=enrichment, run_wgcna=a["run_wgcna"] == "yes",
                orgdb_strategy=a["orgdb_strategy"] if enrichment else "skip",
                orgdb_cache_dir=absolute(str(a["orgdb_cache_dir"])) if a.get("orgdb_cache_dir") else None,
                expected_min_samples=a["min_samples_per_group"], storage_budget_gb=None)
