#!/usr/bin/env python3
"""
generate_intake_template.py
---------------------------------------------------------------------------
Reads config/intake_questions.yaml and writes intake_template.xlsx.

The Excel file has:
  - A README tab explaining how to use it
  - One questionnaire tab per assay (Bulk RNA-seq, TAGseq, sRNA-seq)
  - Each questionnaire tab has the SAME common questions plus assay-specific
    extras at the bottom.

To change the questions, edit config/intake_questions.yaml and re-run:
    python scripts/generate_intake_template.py

That's the whole design. Everything below is plain Python with comments —
poke around freely.
---------------------------------------------------------------------------
"""

from pathlib import Path

import yaml
from openpyxl import Workbook
from openpyxl.styles import (
    Alignment,
    Border,
    Font,
    PatternFill,
    Side,
)
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation


# ===========================================================================
# Paths — where to find the input YAML and where to write the output Excel
# ===========================================================================
REPO_ROOT = Path(__file__).resolve().parent.parent
QUESTIONS_YAML = REPO_ROOT / "config" / "intake_questions.yaml"
OUTPUT_XLSX = REPO_ROOT / "intake_template.xlsx"


# ===========================================================================
# Visual styling — colors and fonts, all in one place so they're easy to tweak
# ===========================================================================

# Yellow background for required-answer cells (so the eye is drawn to them)
REQUIRED_FILL = PatternFill(start_color="FFF9C4", end_color="FFF9C4", fill_type="solid")

# Light gray for optional cells (visually quieter)
OPTIONAL_FILL = PatternFill(start_color="F5F5F5", end_color="F5F5F5", fill_type="solid")

# Dark green for headers — matches the project's overall accent color
HEADER_FILL = PatternFill(start_color="1B4D3E", end_color="1B4D3E", fill_type="solid")

# White text on dark green headers
HEADER_FONT = Font(name="Calibri", size=12, bold=True, color="FFFFFF")

# Bold for the title row at the top of each tab
TITLE_FONT = Font(name="Calibri", size=16, bold=True, color="1B4D3E")

# Italic for the description row beneath the title
DESCRIPTION_FONT = Font(name="Calibri", size=11, italic=True, color="6B6B6B")

# Smaller italic for help text in the rightmost column
HELP_FONT = Font(name="Calibri", size=9, italic=True, color="666666")

# Thin gray border for every cell
THIN_BORDER = Border(
    left=Side(border_style="thin", color="CFD4D0"),
    right=Side(border_style="thin", color="CFD4D0"),
    top=Side(border_style="thin", color="CFD4D0"),
    bottom=Side(border_style="thin", color="CFD4D0"),
)


# ===========================================================================
# Helpers — small utility functions
# ===========================================================================

def load_questions():
    """Read the YAML file and return the parsed dict."""
    with open(QUESTIONS_YAML) as f:
        return yaml.safe_load(f)


def all_questions_for_tab(common_questions, extra_questions):
    """Combine common questions + assay-specific extras into one ordered list."""
    return list(common_questions) + list(extra_questions)


def column_widths():
    """Return a dict of column letter → width.

    Layout:  A=number  B=question  C=answer  D=help
    """
    return {
        "A": 6,    # Q#
        "B": 55,   # question text
        "C": 30,   # answer cell
        "D": 55,   # help / example
    }


# ===========================================================================
# README tab — explains the workflow in plain English at the top of the file
# ===========================================================================

def write_readme_tab(workbook):
    """First tab. The user opens this when they first see the file."""
    ws = workbook.create_sheet(title="README", index=0)

    # Column widths
    for col, width in {"A": 100}.items():
        ws.column_dimensions[col].width = width

    # Write each line into rows. Plain text, one line per row.
    lines = [
        ("This workbook is a PLANNING AID, not an input file", "title"),
        ("", "blank"),
        ("Nothing reads this file. It exists so you can think through, and write", "body"),
        ("down, every decision the pipeline will ask you about, before you start a", "body"),
        ("run. You then type those answers into the prompts.", "body"),
        ("", "blank"),
        ("How to actually start a run", "section"),
        ("1. Fill in the tab that matches your assay, for your own reference.", "body"),
        ("2. From your Mac, upload the metadata spreadsheet and start the wizard:", "body"),
        ("       scripts/start_new.sh <project_name> <metadata_file> [fastq_dir]", "code"),
        ("   That drops you into an interactive session on FARM running", "body"),
        ("       scripts/new_project.sh <project_name>", "code"),
        ("   which auto-detects the sample-ID and factor columns from the", "body"),
        ("   metadata file and prompts you for the rest.", "body"),
        ("3. Read the plan it prints. Type 'y' to launch.", "body"),
        ("", "blank"),
        ("Note: the metadata spreadsheet in step 2 is your OWN sample table", "body"),
        ("(.xlsx/.csv/.tsv, one row per sample). It is not this workbook.", "body"),
        ("", "blank"),
        ("Tips", "section"),
        ("- Use the Help column on the right of each question for examples.", "body"),
        ("- Only fill in ONE tab per project. Leave the other tabs empty.", "body"),
        ("- You can edit the questions themselves: see config/intake_questions.yaml.", "body"),
        ("- If you change the questions, re-run scripts/generate_intake_template.py.", "body"),
        ("", "blank"),
        ("Questions marked 'don't know'", "section"),
        ("Strandedness and library type do not have to be answered by you: Salmon", "body"),
        ("auto-detects the library type during quantification, and the pipeline", "body"),
        ("records what it found and warns you if it disagrees with what you", "body"),
        ("declared in config (samples.expected_libtype).", "body"),
        ("There is no separate probe job; detection happens inside the normal run.", "body"),
    ]

    for row_idx, (text, kind) in enumerate(lines, start=1):
        cell = ws.cell(row=row_idx, column=1, value=text)
        if kind == "title":
            cell.font = TITLE_FONT
        elif kind == "section":
            cell.font = Font(name="Calibri", size=13, bold=True, color="1B4D3E")
        elif kind == "code":
            cell.font = Font(name="Courier New", size=10, color="333333")
        else:
            cell.font = Font(name="Calibri", size=11)
        cell.alignment = Alignment(vertical="top", wrap_text=True)


# ===========================================================================
# Questionnaire tab — one per assay
# ===========================================================================

def write_questionnaire_tab(workbook, tab_name, tab_description, questions):
    """Write a single questionnaire tab.

    Each tab has the same layout:
      Row 1:  tab title
      Row 2:  tab description
      Row 3:  blank
      Row 4:  column headers (Q#, Question, Your answer, Help / example)
      Row 5+: one row per question
    """
    ws = workbook.create_sheet(title=tab_name)

    # --- Set column widths ---
    for col, width in column_widths().items():
        ws.column_dimensions[col].width = width

    # --- Row 1: title ---
    ws.cell(row=1, column=1, value=tab_name).font = TITLE_FONT
    ws.merge_cells(start_row=1, end_row=1, start_column=1, end_column=4)

    # --- Row 2: description ---
    description_cell = ws.cell(row=2, column=1, value=tab_description.strip())
    description_cell.font = DESCRIPTION_FONT
    description_cell.alignment = Alignment(vertical="top", wrap_text=True)
    ws.merge_cells(start_row=2, end_row=2, start_column=1, end_column=4)
    ws.row_dimensions[2].height = 50

    # --- Row 4: column headers ---
    headers = ["Q#", "Question", "Your answer", "Help / example"]
    for col_idx, header_text in enumerate(headers, start=1):
        cell = ws.cell(row=4, column=col_idx, value=header_text)
        cell.font = HEADER_FONT
        cell.fill = HEADER_FILL
        cell.alignment = Alignment(horizontal="left", vertical="center")
        cell.border = THIN_BORDER

    # --- Rows 5+: one per question ---
    for q_idx, question in enumerate(questions, start=1):
        row = q_idx + 4  # questions start at row 5

        # Column A: question number
        cell_num = ws.cell(row=row, column=1, value=q_idx)
        cell_num.alignment = Alignment(horizontal="center", vertical="top")
        cell_num.border = THIN_BORDER

        # Column B: question text
        cell_text = ws.cell(row=row, column=2, value=question["text"])
        cell_text.alignment = Alignment(vertical="top", wrap_text=True)
        cell_text.border = THIN_BORDER

        # Column C: answer cell (with pre-fill, validation, color)
        answer_value = question.get("default", "")
        cell_answer = ws.cell(row=row, column=3, value=answer_value)
        cell_answer.alignment = Alignment(vertical="top", wrap_text=True)
        cell_answer.border = THIN_BORDER

        if question.get("required", False):
            cell_answer.fill = REQUIRED_FILL
        else:
            cell_answer.fill = OPTIONAL_FILL

        # If this question is a dropdown, attach a data-validation list to
        # the answer cell. Excel will then show a dropdown arrow.
        if question["type"] == "dropdown":
            # Quotes are required around comma-separated values for openpyxl
            options = ",".join(question["options"])
            dv = DataValidation(
                type="list",
                formula1=f'"{options}"',
                allow_blank=True,
                showDropDown=False,   # arrow visible (Excel quirk: False = visible)
            )
            ws.add_data_validation(dv)
            dv.add(cell_answer)

        # Column D: help / example
        cell_help = ws.cell(row=row, column=4, value=question.get("help", ""))
        cell_help.font = HELP_FONT
        cell_help.alignment = Alignment(vertical="top", wrap_text=True)
        cell_help.border = THIN_BORDER

        # Row height: enough for ~2 lines of wrapped text
        ws.row_dimensions[row].height = 35

    # Freeze the top 4 rows so they stay visible when scrolling
    ws.freeze_panes = "A5"


# ===========================================================================
# Main — orchestrate the whole thing
# ===========================================================================

def main():
    print(f"Reading questions from: {QUESTIONS_YAML}")
    config = load_questions()

    common = config.get("common_questions", [])
    tabs = config.get("tabs", [])

    print(f"  - {len(common)} common questions")
    print(f"  - {len(tabs)} assay tabs: {[t['name'] for t in tabs]}")

    # Create a fresh workbook (openpyxl gives us one default sheet; we'll
    # remove it and create our own named tabs).
    wb = Workbook()
    default_sheet = wb.active
    wb.remove(default_sheet)

    # README first (becomes the leftmost tab)
    write_readme_tab(wb)

    # Then one questionnaire tab per assay
    for tab_def in tabs:
        tab_name = tab_def["name"]
        tab_description = tab_def.get("description", "")
        extras = tab_def.get("extra_questions", [])
        all_questions = all_questions_for_tab(common, extras)

        print(f"  Writing tab '{tab_name}' ({len(all_questions)} questions)")
        write_questionnaire_tab(wb, tab_name, tab_description, all_questions)

    # Save
    wb.save(OUTPUT_XLSX)
    print(f"\nWrote: {OUTPUT_XLSX}")
    print("\nThis workbook is a planning aid: nothing parses it.")
    print("Fill in the tab for your assay to decide your answers, then start a run with")
    print("  scripts/start_new.sh <project_name> <metadata_file> [fastq_dir]")


if __name__ == "__main__":
    main()
